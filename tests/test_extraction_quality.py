"""Extraction quality — the hand test's plain fixes (spec §8.7 "Extraction quality", D-040).

F-06 Word text boxes, F-07 Office Math, F-08 empty slides and pages, F-03/F-09/F-10 library noise
and ligatures, F-12 titles. Every source is generated here, in `tmp_path` — a DOCX with text boxes
inside `mc:AlternateContent` and an OMML equation are written as raw XML through python-docx and
python-pptx. No real course material, no binary fixtures, no network.

What cannot be tested here: whether fontTools really repairs the broken characters of a real
textbook (F-09). That is for the teacher's re-test on the CLRS PDF.
"""

from __future__ import annotations

import logging
import warnings
from pathlib import Path

import pytest
from materials_fixtures import make_docx, make_pdf, make_pptx

from classkit import ingest
from classkit.cli import main
from classkit.ingest import extract
from classkit.ingest.manifest import load
from classkit.scaffold import scaffold_course

FRAMEWORK_ROOT = Path(__file__).resolve().parents[1]

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(*_args, **_kwargs):
        raise OSError("no network in tests")

    monkeypatch.setattr("urllib.request.urlopen", refuse)


@pytest.fixture
def course(tmp_path: Path) -> Path:
    root = tmp_path / "course"
    scaffold_course(root, FRAMEWORK_ROOT, code="T-1", title="Test", institution="U",
                    instructor="", units=13, methodology="question-driven-25")
    return root


def source(course: Path) -> Path:
    return course / "materials" / "source"


def text_of(course: Path, material_id: str = "M0001") -> str:
    return ingest.ingested_file(course, material_id).read_text(encoding="utf-8")


# -- raw OOXML, written through python-docx / python-pptx ----------------------------------

TEXT_BOX = f"""
<w:r xmlns:w="{W}" xmlns:mc="{MC}"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="{A}" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx>
        <w:txbxContent>
          <w:p><w:r><w:t>{{text}}</w:t></w:r></w:p>
        </w:txbxContent>
      </wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict><v:shape><v:textbox>
        <w:txbxContent>
          <w:p><w:r><w:t>{{text}}</w:t></w:r></w:p>
        </w:txbxContent>
      </v:textbox></v:shape></w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>"""

FRACTION = f"""
<m:oMath xmlns:m="{M}">
  <m:f><m:num><m:r><m:t>n</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>
  <m:r><m:t>+</m:t></m:r>
  <m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>
</m:oMath>"""


def docx_with(path: Path, *, boxes: list[str] = (), equation: bool = False, filler: int = 0) -> Path:
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    document.add_paragraph("Syllabus")
    for text in boxes:
        anchor = document.add_paragraph("e")  # what the hand test's syllabus showed: one letter
        anchor._p.append(parse_xml(TEXT_BOX.replace("{text}", text)))
    if equation:
        paragraph = document.add_paragraph("The cost is ")
        paragraph._p.append(parse_xml(FRACTION))
    for _ in range(filler):  # heavy XML, no text: what pictures and drawings look like
        paragraph = document.add_paragraph()
        run = paragraph.add_run("")
        run.bold, run.italic, run.underline = True, True, True
        run.font.name = "Times New Roman"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


# -- F-06: text boxes ------------------------------------------------------------------

def test_docx_text_boxes_are_extracted_once(course: Path):
    docx_with(source(course) / "syllabus.docx", boxes=["Credits: 5 ECTS", "Lecturer: Prof. A"])
    ingest.run(course, fetch=False)
    text = text_of(course)
    assert text.count("Credits: 5 ECTS") == 1  # the Choice or the Fallback, never both
    assert text.count("Lecturer: Prof. A") == 1
    assert text.index("Syllabus") < text.index("Credits") < text.index("Lecturer")


def test_text_boxes_from_the_fallback_when_there_is_no_choice(course: Path):
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    xml = TEXT_BOX.replace("{text}", "Only in VML")
    start, end = xml.index("<mc:Choice"), xml.index("</mc:Choice>") + len("</mc:Choice>")
    document.add_paragraph("x")._p.append(parse_xml(xml[:start] + xml[end:]))
    path = source(course) / "old.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    ingest.run(course, fetch=False)
    assert text_of(course).count("Only in VML") == 1


# -- F-07: equations -------------------------------------------------------------------

def test_a_docx_equation_is_linear_text():
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    paragraph = document.add_paragraph("The cost is ")
    paragraph._p.append(parse_xml(FRACTION))
    path = Path(__file__)  # unused: extraction goes through the element
    del path
    text = extract._docx_paragraph_text(paragraph._p, paragraph)
    assert "`(n)/(2)+x^(2)`" in text


def test_docx_equation_reaches_the_ingested_file(course: Path):
    docx_with(source(course) / "notes.docx", equation=True)
    ingest.run(course, fetch=False)
    assert "The cost is `(n)/(2)+x^(2)`" in text_of(course)


def test_a_slide_equation_is_linear_text(course: Path):
    from lxml import etree
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Loop invariant"
    paragraph = slide.placeholders[1].text_frame.paragraphs[0]
    paragraph.text = "Invariant:"
    alternate = etree.fromstring(f"""
<mc:AlternateContent xmlns:mc="{MC}" xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">
  <mc:Choice Requires="a14"><a14:m><m:oMathPara xmlns:m="{M}">{FRACTION}</m:oMathPara></a14:m></mc:Choice>
  <mc:Fallback><a:r xmlns:a="{A}"><a:t></a:t></a:r></mc:Fallback>
</mc:AlternateContent>""")
    paragraph._p.append(alternate)
    path = source(course) / "unit2.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))

    ingest.run(course, fetch=False)
    assert "Invariant: `(n)/(2)+x^(2)`" in text_of(course)


def test_an_unreadable_equation_leaves_a_placeholder():
    from lxml import etree

    empty = etree.fromstring(f'<m:oMath xmlns:m="{M}"><m:r><m:t></m:t></m:r></m:oMath>')
    assert extract.equation(empty) == "[equation]"


# -- F-06 / F-08: thin extraction is reported -------------------------------------------

def test_a_low_yield_document_is_flagged_in_the_summary(course: Path, capsys):
    docx_with(source(course) / "syllabus.docx", filler=1500)
    main(["ingest", "--no-fetch", "--course", str(course)])
    out = capsys.readouterr().out
    assert "thin or degraded text is not thin teaching" in out
    assert "only 8 characters of text from a" in out and "M0001" in out


def test_an_ordinary_document_is_not_flagged(course: Path, capsys):
    make_docx(source(course) / "notes.docx", [("h1", "Heaps"), ("p", "A heap is a tree. " * 400)])
    main(["ingest", "--no-fetch", "--course", str(course)])
    assert "Check these extractions" not in capsys.readouterr().out


def test_empty_slides_are_counted(course: Path, capsys):
    from pptx import Presentation

    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Heaps"
    deck.slides.add_slide(deck.slide_layouts[6])  # blank: a picture-only slide
    deck.slides.add_slide(deck.slide_layouts[6])
    path = source(course) / "deck.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))
    main(["ingest", "--no-fetch", "--course", str(course)])
    assert "2 of 3 slides have no text" in capsys.readouterr().out


def test_empty_pages_are_counted(course: Path, capsys):
    make_pdf(source(course) / "chapter.pdf", ["Heaps are complete binary trees", "",
                                              "Heapsort runs in n log n time"])
    main(["ingest", "--no-fetch", "--course", str(course)])
    assert "1 of 3 pages have no text" in capsys.readouterr().out


# -- F-03 / F-10: library noise is captured -------------------------------------------------

def test_library_warnings_are_captured_and_reported_once_by_name(course: Path, monkeypatch, capfd):
    real = extract.EXTRACTORS[".pdf"].extract

    def noisy(path):
        for _ in range(180):
            logging.getLogger("pypdf._cmap").warning(
                "fontTools is required to fully parse the encoding of a CFF Type1 font {'/Font': …}")
        warnings.warn("Previous trailer cannot be read")
        return real(path)

    monkeypatch.setattr(extract.EXTRACTORS[".pdf"], "extract", noisy)
    make_pdf(source(course) / "book.pdf", ["Heaps are complete binary trees"])

    main(["ingest", "--no-fetch", "--course", str(course)])

    captured = capfd.readouterr()
    assert "fontTools is required" not in captured.err
    assert captured.out.count("fontTools is required") == 1  # once, as the example
    assert "181 problem(s)" in captured.out and "M0001" in captured.out


def test_a_malformed_pdf_says_nothing_on_the_terminal(course: Path, capfd):
    path = make_pdf(source(course) / "broken-xref.pdf", ["Heaps", "Heapsort"])
    data = path.read_bytes()
    position = data.rindex(b"startxref\n") + len(b"startxref\n")
    path.write_bytes(data[:position] + b"999" + data[data.index(b"\n", position):])

    main(["ingest", "--preflight", "--no-fetch", "--course", str(course)])
    main(["ingest", "--no-fetch", "--course", str(course)])

    captured = capfd.readouterr()
    assert captured.err == ""
    assert "incorrect startxref pointer" in captured.out  # reported, by name, in the reports
    assert "## Page 2" in text_of(course)


def test_the_reader_dependency_for_fonts_is_installed():
    """F-09's likely cause: pypdf parses CFF/Type1 fonts only with fontTools."""
    import fontTools  # noqa: F401


# -- F-09: ligatures --------------------------------------------------------------------

def test_pdf_ligatures_are_expanded(course: Path, monkeypatch):
    from pypdf import PageObject

    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: "eﬃcient ﬁrst ﬂow oﬀer")
    make_pdf(source(course) / "book.pdf", ["placeholder"])
    ingest.run(course, fetch=False)
    assert "efficient first flow offer" in text_of(course)


# -- F-12: titles ----------------------------------------------------------------------

def pdf_with_metadata(path: Path, title: str, pages: list[str]) -> Path:
    from pypdf import PdfWriter

    make_pdf(path, pages)
    writer = PdfWriter(clone_from=str(path))
    writer.add_metadata({"/Title": title})
    with path.open("wb") as stream:
        writer.write(stream)
    return path


@pytest.mark.parametrize("junk", ["manual.dvi", "chapter3.tex", "Microsoft Word - syllabus.docx",
                                  "C:\\Users\\x\\notes.doc"])
def test_a_metadata_title_that_is_a_file_name_is_rejected(course: Path, junk: str):
    pdf_with_metadata(source(course) / "Cormen_4e_Instructors_Manual.pdf", junk, ["Copyright 2022"])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "Cormen 4e Instructors Manual"


def test_a_real_metadata_title_is_kept(course: Path):
    pdf_with_metadata(source(course) / "x.pdf", "Introduction to Algorithms", ["Copyright 2022"])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "Introduction to Algorithms"


def test_a_pdfs_first_line_is_never_its_title(course: Path):
    make_pdf(source(course) / "lecture_notes_week_3.pdf", ["Copyright 2022 The MIT Press"])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "lecture notes week 3"


# -- F-14, F-20 -------------------------------------------------------------------------

def test_the_estimate_for_a_textbook_sized_run_is_honest():
    """Teacher test: 2,204 PDF pages took ~30 s on a recent Mac; the estimate said "about 3
    minutes" (F-14 had been calibrated on a slower machine)."""
    from classkit.ingest.core import SECONDS_PER_PAGE
    from classkit.ingest.report import _duration

    assert 20 <= 2204 * SECONDS_PER_PAGE <= 60
    assert _duration(2204 * SECONDS_PER_PAGE) == "under a minute"


def test_the_session_template_says_the_budget_check_warns():
    template = (FRAMEWORK_ROOT / "templates" / "unit" / "session.md").read_text(encoding="utf-8")
    assert "validator fails" not in template and "validator warns" in template


# -- F-23, F-25 -------------------------------------------------------------------------

def test_preflight_names_what_changed(course: Path, capsys):
    deck = make_pptx(source(course) / "unit1.pptx", [("Heaps", "a")])
    make_pdf(source(course) / "old.pdf", ["x"])
    ingest.run(course, fetch=False)
    make_pptx(deck, [("Heaps", "a, revised")])
    (source(course) / "old.pdf").rename(source(course) / "renamed.pdf")
    make_pdf(source(course) / "new.pdf", ["y"])

    main(["ingest", "--preflight", "--no-fetch", "--course", str(course)])

    out = capsys.readouterr().out
    ids = {r["canonical"]: r["id"] for r in load(course)}
    assert f"changed  {ids['unit1.pptx']}: unit1.pptx" in out
    assert f"moved    {ids['old.pdf']}: old.pdf → renamed.pdf" in out
    assert "new      new.pdf" in out


def test_after_a_refusal_validate_names_the_decision_not_run_ingest(course: Path, capsys):
    from classkit.model import load_course
    from classkit.validate import validate

    deck = make_pptx(source(course) / "unit1.pptx", [("Heaps", "a")])
    ingest.run(course, fetch=False)
    path = ingest.ingested_file(course, "M0001")
    path.write_text(path.read_text(encoding="utf-8") + "\nfixed by hand\n", encoding="utf-8")
    make_pptx(deck, [("Heaps", "a, revised")])
    assert main(["ingest", "--no-fetch", "--course", str(course)]) == 3

    (found,) = [f for f in validate(load_course(course, FRAMEWORK_ROOT), FRAMEWORK_ROOT)
                if f.code == "materials_not_ingested"]
    assert "classkit ingest --keep M0001" in found.message and "--overwrite M0001" in found.message
    assert "Run /ingest" not in found.message


def test_the_links_template_does_not_promise_harvesting():
    """D-040: links inside slides are not collected; the template said they were (teacher test)."""
    template = (FRAMEWORK_ROOT / "templates" / "course" / "links.md").read_text(encoding="utf-8")
    assert "collected automatically" not in template
    assert "NOT\ncollected" in template or "not collected" in template.lower().replace("\n", " ")


# -- deck titles: slide 1 or the file name (teacher test, F-12 follow-up) -------------------

def test_a_deck_title_comes_from_slide_1(course: Path):
    make_pptx(source(course) / "Unit_2_Heaps.pptx", [("Heaps and heapsort", "Today"), ("Max-heapify", "x")])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "Heaps and heapsort"


def test_a_deck_title_never_comes_from_a_later_slide(course: Path):
    """The teacher's Unit 4 deck was titled from slide 2, a tribute slide."""
    make_pptx(source(course) / "Unit_4_Divide-and-Conquer_board.pptx",
              [("", "Last Week"), ("In Memoriam: Michael O. Rabin", "1931–2026")])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "Unit 4 Divide and Conquer board"


def test_a_deck_with_a_boilerplate_slide_1_title_uses_its_file_name(course: Path):
    make_pptx(source(course) / "unit_3.pptx", [("PowerPoint Presentation", "x")])
    ingest.run(course, fetch=False)
    assert load(course)[0]["title"] == "unit 3"


@pytest.mark.parametrize("name, title", [
    ("Introduction.to.Algorithms.4th.Edition.2022.4.pdf", "Introduction to Algorithms 4th Edition 2022.4"),
    ("Lecture-Notes-for-Chapter-1.pdf", "Lecture Notes for Chapter 1"),
    ("notes_ch6.2.pdf", "notes ch6.2"),
])
def test_a_title_from_a_file_name_is_readable(name: str, title: str):
    assert extract.title_from_name(name) == title


# -- ligature damage: the `û` repair and the garbled-text probe (teacher test, CLRS 4e) ------

def test_a_fi_substitute_is_repaired_inside_an_ascii_word():
    assert extract.repair_ligatures("an efûcient way to ûnd the Ûrst") == "an efficient way to find the Ûrst"


def test_a_real_accented_word_keeps_its_letter():
    assert extract.repair_ligatures("crème brûlée and Brûlé") == "crème brûlée and Brûlé"


def test_the_pdf_text_is_repaired_on_ingest(course: Path, monkeypatch):
    from pypdf import PageObject

    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: "an efûcient way to ûnd it " * 50)
    make_pdf(source(course) / "book.pdf", ["placeholder"])
    ingest.run(course, fetch=False)
    assert "efficient way to find" in text_of(course) and "û" not in text_of(course)


def test_the_probe_counts_words_with_one_stray_latin_letter():
    text = "an eÿcient algorithm " * 30 + "plain words " * 100
    assert extract.garbled(text) == (30, "eÿcient")


def test_the_probe_ignores_a_few_names_and_greek_mathematics():
    text = "Erdős and the naïve bound " + "Θn and πr and Θlgn " * 40 + "plain words " * 2000
    assert extract.garbled(text) is None


def test_the_probe_needs_a_material_share_not_only_a_count():
    text = "an eÿcient algorithm " * 25 + "plain words here " * 3000
    assert extract.garbled(text) is None  # 25 of ~9,075 words: under 0.5%


def test_garbled_text_is_reported_in_check_these_extractions(course: Path, monkeypatch, capsys):
    from pypdf import PageObject

    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: "an eÿcient sort " * 40)
    make_pdf(source(course) / "book.pdf", ["placeholder"])
    main(["ingest", "--no-fetch", "--course", str(course)])
    out = capsys.readouterr().out
    assert "Check these extractions" in out
    assert 'M0001  ~40 words look garbled (e.g. "eÿcient") — a font-encoding problem' in out


def test_repaired_text_is_not_reported(course: Path, monkeypatch, capsys):
    from pypdf import PageObject

    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: "an efûcient sort " * 40)
    make_pdf(source(course) / "book.pdf", ["placeholder"])
    main(["ingest", "--no-fetch", "--course", str(course)])
    assert "garbled" not in capsys.readouterr().out


def test_preflight_reader_warnings_only_for_files_about_to_be_converted(course: Path, capsys):
    """Teacher test: an unchanged file's reader warnings were repeated on every pre-flight."""
    path = make_pdf(source(course) / "broken-xref.pdf", ["Heaps", "Heapsort"])
    data = path.read_bytes()
    position = data.rindex(b"startxref\n") + len(b"startxref\n")
    path.write_bytes(data[:position] + b"999" + data[data.index(b"\n", position):])

    main(["ingest", "--preflight", "--no-fetch", "--course", str(course)])
    first = capsys.readouterr().out
    assert "incorrect startxref pointer" in first
    assert "Usually harmless; check these files' extraction after ingest." in first

    main(["ingest", "--no-fetch", "--course", str(course)])
    capsys.readouterr()
    make_pdf(source(course) / "new.pdf", ["Quicksort"])
    main(["ingest", "--preflight", "--no-fetch", "--course", str(course)])
    again = capsys.readouterr().out
    assert "incorrect startxref pointer" not in again and "Read with warnings" not in again


# -- DOCX: a bold label is a heading (teacher test: the BGU syllabus form) ------------------

def docx_with_labels(path: Path) -> Path:
    import docx

    document = docx.Document()
    def para(*runs):
        paragraph = document.add_paragraph()
        for text, bold in runs:
            paragraph.add_run(text).bold = bold
    para(("Course syllabus", False))
    para(("Assessment:", True))
    para(("Exam 70%, homework 30%.", False))
    para(("Learning outcomes ", True), ("of the module:", True))
    para(("Explain heaps.", False))
    para(("Not a label: ", True), ("the rest is plain", False))
    para(("Bold but no colon", True))
    para(("A bold sentence that is far too long to be a section label in any form at all, in any language at all:", True))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def test_a_short_bold_label_ending_in_a_colon_is_an_anchor(course: Path):
    docx_with_labels(source(course) / "syllabus.docx")
    ingest.run(course, fetch=False)
    text = text_of(course)
    assert extract.anchors(text) == ["assessment", "learning-outcomes-of-the-module"]
    assert "## Assessment\n" in text
    assert load(course)[0]["title"] == "syllabus"  # a label is not the document's title
