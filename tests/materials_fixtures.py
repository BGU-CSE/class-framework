"""Small source files built on the fly, for the ingest tests.

No real course material and no binary sample files are committed (invariant 1): every deck,
PDF and Word file a test reads is generated here, into the test's `tmp_path`.
"""

from __future__ import annotations

from pathlib import Path


def make_pptx(path: Path, slides: list[tuple[str, str]], *, notes: dict[int, str] | None = None,
              link: tuple[int, str] | None = None) -> Path:
    """A deck with one (title, body) per slide. `link` = (slide number, URL) hyperlinked."""
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[1]  # title and content
    for number, (title, body) in enumerate(slides, start=1):
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = body
        if link and link[0] == number:
            run = frame.paragraphs[0].add_run()
            run.text = " (more)"
            run.hyperlink.address = link[1]
        if notes and number in notes:
            slide.notes_slide.notes_text_frame.text = notes[number]
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))
    return path


def make_docx(path: Path, blocks: list[tuple[str, str]], *, link: str | None = None) -> Path:
    """A Word file from (style, text) blocks: style 'h1', 'h2', or 'p'."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    for style, text in blocks:
        if style.startswith("h"):
            document.add_heading(text, level=int(style[1:]))
        else:
            document.add_paragraph(text)
    if link:
        paragraph = document.add_paragraph("See ")
        rel = paragraph.part.relate_to(link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "the reference"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def make_pdf(path: Path, pages: list[str]) -> Path:
    """A minimal PDF with one line of Helvetica text per page ("" makes a blank page — a
    stand-in for a scan with no text layer). Built by hand so the tests need no PDF writer."""
    objects: list[bytes] = []
    kids = []
    font_id = 3 + 2 * len(pages)
    for index, text in enumerate(pages):
        page_id, content_id = 3 + 2 * index, 4 + 2 * index
        kids.append(f"{page_id} 0 R")
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1") if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    catalog = b"<< /Type /Catalog /Pages 2 0 R >>"
    pages_obj = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    font = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    ordered = [catalog, pages_obj, *objects, font]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(ordered, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(ordered) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(ordered) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))
    return path


def make_book(path: Path, pages: list[str], *, outline: list[tuple[str, int]] = (),
              front_matter: int = 0) -> Path:
    """A book-like PDF: `make_pdf`'s pages, plus an outline (bookmarks: (title, page index)) and
    printed page labels — the first `front_matter` pages numbered i, ii, …, the rest 1, 2, … —
    so the physical page and the printed one differ, as in a real textbook. Written with pypdf."""
    from pypdf import PdfWriter

    make_pdf(path, pages)
    writer = PdfWriter(clone_from=str(path))
    for title, page in outline:
        writer.add_outline_item(title, page)
    if front_matter:
        writer.set_page_label(0, front_matter - 1, style="/r")
        writer.set_page_label(front_matter, len(pages) - 1, style="/D", start=1)
    with path.open("wb") as stream:
        writer.write(stream)
    return path


def make_annotated_pdf(path: Path, pages: list[str]) -> Path:
    """`make_pdf`'s pages, annotated as a teacher marks lecture notes (D-048): on page 1 a
    highlight over the line of text (with a note), a sticky note, a free-text box and a stamp;
    the other pages unannotated. The text sits at (72, 720), 12 pt, so the highlight covers it."""
    from pypdf import PdfWriter
    from pypdf.annotations import FreeText, Highlight, Text
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    make_pdf(path, pages)
    writer = PdfWriter(clone_from=str(path))
    quads = ArrayObject([FloatObject(v) for v in (70, 734, 400, 734, 70, 716, 400, 716)])
    highlight = Highlight(rect=(70, 716, 400, 734), quad_points=quads)
    highlight[NameObject("/Contents")] = TextStringObject("prove this on the board")
    writer.add_annotation(0, highlight)
    writer.add_annotation(0, Text(rect=(450, 700, 470, 720), text="Ask in class what the running time is"))
    writer.add_annotation(0, FreeText(text="Show on cards", rect=(72, 600, 300, 630)))
    stamp = DictionaryObject({NameObject("/Type"): NameObject("/Annot"),
                              NameObject("/Subtype"): NameObject("/Stamp"),
                              NameObject("/Rect"): ArrayObject([FloatObject(v) for v in (300, 500, 400, 540)]),
                              NameObject("/Name"): NameObject("/Approved")})
    writer.add_annotation(0, stamp)
    with path.open("wb") as stream:
        writer.write(stream)
    return path
