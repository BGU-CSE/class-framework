"""Extractors: one source file in, Markdown with explicit anchors out (spec §8.7).

Conversion is code, not agent work (D-018): it must produce **the same anchors every time**,
or locators rot. Each format's anchor:

    PPTX                    one `## Slide N` per slide, N = position in the deck
    PDF                     one `## Page N` per page, N = physical page (1-based)
    DOCX, MD, and pandoc    the document's own headings
    TXT                     none — a plain-text file has no structure to point into

Adding a format is one function registered by extension (`@register`) — an extension point
like adding a methodology. Anything without an extractor is recorded `unsupported` with a hint,
never dropped; audio and video are `media`; a PDF with no text layer is `no-text`.

The third-party readers (python-pptx, python-docx, pypdf) are imported inside the functions
that use them, so `classkit validate` — which imports this package to compare hashes — does
not pay for them.
"""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
import tempfile
import unicodedata
import warnings
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

INGESTED = "ingested"
UNSUPPORTED = "unsupported"
NO_TEXT = "no-text"
MEDIA = "media"
LINK = "link"

#: Below this many non-blank characters per page on average, a PDF is treated as having no
#: text layer (a scan). A scan often carries a few stray characters — page numbers, a stamp.
NO_TEXT_CHARS_PER_PAGE = 10

#: A zip-based document (DOCX, PPTX) whose text is under this many characters per byte of its
#: document XML came out "far smaller than its source" — text boxes, pictures, objects not read.
#: Ordinary documents run 0.02–0.1; the hand test's text-box syllabus was 13 characters (F-06).
LOW_YIELD_RATIO = 0.005
#: Below this much document XML, a short document is simply short.
LOW_YIELD_MIN_XML = 20_000

#: The third-party loggers whose warnings are captured, never printed (F-03, F-10).
NOISY_LOGGERS = ("pypdf", "PyPDF2", "pptx", "docx", "fontTools")

#: Typographic ligatures a PDF's text layer often carries as single characters (F-09).
LIGATURES = str.maketrans({"\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi",
                           "\ufb04": "ffl", "\ufb05": "st", "\ufb06": "st"})

#: A letter run — a word, for the ligature repair and the garbled-text probe.
_WORD = re.compile(r"[^\W\d_]+")
#: Some PDFs' fonts map the "fi" ligature to `û` (CLRS 4e: `efûcient`, `ûnd` — ~4,700 words),
#: and fontTools does not help (teacher test). Repaired only inside a word whose other letters
#: are all ASCII, so a real `û` beside other accented letters (a French name) survives.
_FI_SUBSTITUTE = "\u00fb"

#: The garbled-text probe (teacher test): a material is reported when at least this many words…
GARBLED_MIN_WORDS = 20
#: …and at least this share of its words are ASCII but for one Latin non-ASCII letter (`efûcient`).
#: CLRS 4e unrepaired ran ~1%; ordinary English text with a few names (`Erdős`, `naïve`) runs far below.
GARBLED_MIN_SHARE = 0.005

MEDIA_EXTENSIONS = {
    ".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".wmv",
    ".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac",
}


@dataclass
class Extraction:
    """What an extractor produced from one source file."""

    status: str
    body: str = ""
    title: str = ""
    reason: str = ""
    #: anchor → one-line labels for a private material's index (D-040): a slide's title, a
    #: page's printed label and the sections that start on it. Never body text.
    labels: dict[str, list[str]] = field(default_factory=dict)
    #: True when `title` was taken from the body text (a text file's first line) rather than a
    #: heading, a slide title or metadata — so a private material does not carry a line of its
    #: text into the committed manifest and index. A PDF's first line is never a title (F-12).
    title_from_body: bool = False
    #: what the third-party reader complained about while reading it — captured, not printed
    warnings: list[str] = field(default_factory=list)
    #: (empty, total, "slides" | "pages") — how many slides or pages came out with no text (F-08)
    empty: tuple[int, int, str] | None = None
    #: set when the text is far smaller than the source (F-06): what to tell the teacher
    low_yield: str = ""
    #: (count, example) — words that look garbled by a font-encoding problem; see `garbled()`
    garbled: tuple[int, str] | None = None

    def quality(self) -> list[str]:
        """One line per thing the teacher should know about this extraction — low yield, empty
        slides or pages, a reader that complained — so thin extraction is not mistaken for thin
        teaching. Empty when there is nothing to say."""
        notes = []
        if self.low_yield:
            notes.append(self.low_yield)
        if self.empty and self.empty[0]:
            empty, total, unit = self.empty
            notes.append(f"{empty} of {total} {unit} have no text (pictures? OCR is not done in Core)")
        if self.garbled:
            count, example = self.garbled
            notes.append(f"~{count:,} words look garbled (e.g. \"{example}\") — a font-encoding "
                         "problem; check the extraction (ignore this if the material is in a "
                         "language with accented letters)")
        if self.warnings:
            first = self.warnings[0]
            notes.append(f"the reader reported {len(self.warnings)} problem(s), e.g. \"{first}\" — "
                         "the text may be degraded")
        return notes


@dataclass
class Probe:
    """What the pre-flight report can learn about a file without converting it."""

    slides: int = 0
    pages: int = 0
    status: str = INGESTED  # the status conversion is expected to produce
    reason: str = ""
    #: what the reader complained about while probing — captured, reported once by name
    warnings: list[str] = field(default_factory=list)


@dataclass
class Extractor:
    extract: Callable[[Path], Extraction]
    probe: Callable[[Path], Probe]
    #: rough seconds per file, for the pre-flight time estimate
    seconds: float = 0.2


#: The registry: file extension (lowercase, with the dot) → extractor.
EXTRACTORS: dict[str, Extractor] = {}


def register(*extensions: str, probe: Callable[[Path], Probe] | None = None, seconds: float = 0.2):
    """Register an extractor for one or more extensions — how a format is added."""

    def decorate(function: Callable[[Path], Extraction]):
        for extension in extensions:
            EXTRACTORS[extension.lower()] = Extractor(function, probe or (lambda _p: Probe()), seconds)
        return function

    return decorate


def format_of(path: Path) -> str:
    return path.suffix.lower().lstrip(".") or "none"


def title_from_name(path: Path | str) -> str:
    """A title from a file name: the stem, with `_`, `-` and `.` read as spaces —
    `Unit_2_Heaps.pptx` → `Unit 2 Heaps`, `Introduction.to.Algorithms.pdf` → `Introduction to
    Algorithms`. A dot between two digits stays (`ch6.2` → `ch6.2`): it is a section number."""
    stem = Path(path).stem
    readable = re.sub(r"(?<!\d)\.|\.(?!\d)", " ", stem.replace("_", " ").replace("-", " "))
    return _clean(readable) or stem


class _Collect(logging.Handler):
    def __init__(self, messages: list[str]):
        super().__init__(logging.WARNING)
        self.messages = messages

    def emit(self, record: logging.LogRecord) -> None:
        self.messages.append(_one_line(record.getMessage()))


def _one_line(message: str, limit: int = 120) -> str:
    """A library's message, first line only and short — pypdf's can dump a whole font dictionary."""
    text = _clean(str(message).splitlines()[0] if str(message).strip() else "")
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


@contextmanager
def quiet(messages: list[str]):
    """Capture what the third-party readers log or warn while reading a file, instead of letting
    it reach the teacher's terminal (F-03, F-10): 180 lines of font warnings buried the summary.
    What was captured is reported once, by file name, in the summary."""
    handler = _Collect(messages)
    saved = []
    for name in NOISY_LOGGERS:
        logger = logging.getLogger(name)
        saved.append((logger, logger.propagate))
        logger.addHandler(handler)
        logger.propagate = False
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            yield
        messages += [_one_line(str(w.message)) for w in caught]
    finally:
        for logger, propagate in saved:
            logger.removeHandler(handler)
            logger.propagate = propagate


def extract(path: Path) -> Extraction:
    """Convert one file. Never raises for a bad file: an unreadable one is `unsupported`."""
    extension = path.suffix.lower()
    if extension in MEDIA_EXTENSIONS:
        return Extraction(MEDIA, title=title_from_name(path),
                          reason="audio/video: recorded, content not extracted in Core")
    extractor = EXTRACTORS.get(extension)
    if extractor is None:
        return Extraction(UNSUPPORTED, title=title_from_name(path), reason=_unsupported_hint(extension))
    captured: list[str] = []
    try:
        with quiet(captured):
            result = extractor.extract(path)
    except Exception as exc:  # a corrupt or password-protected file must not stop the run
        return Extraction(UNSUPPORTED, title=title_from_name(path),
                          reason=f"could not be read ({type(exc).__name__}: {_one_line(str(exc))}); "
                                 "try exporting it to PDF", warnings=captured)
    result.warnings = result.warnings + captured
    result.title = _clean(result.title) or title_from_name(path)
    if result.status == INGESTED:
        result.garbled = garbled(result.body)
    return result


def repair_ligatures(text: str) -> str:
    """`û` → `fi` inside a word whose other letters are ASCII (`efûcient` → `efficient`); a word
    with another non-ASCII letter is left alone. See `_FI_SUBSTITUTE`."""
    if _FI_SUBSTITUTE not in text:
        return text

    def fix(match: re.Match) -> str:
        word = match.group(0)
        if _FI_SUBSTITUTE not in word:
            return word
        rest = word.replace(_FI_SUBSTITUTE, "")
        if rest and rest.isascii():
            return word.replace(_FI_SUBSTITUTE, "fi")
        return word

    return _WORD.sub(fix, text)


def _looks_garbled(word: str) -> bool:
    """ASCII letters but for exactly one Latin non-ASCII letter, with at least two ASCII letters
    beside it: `efûcient`. Greek (Θ, π — mathematics) and other scripts are not counted."""
    if word.isascii():
        return False
    foreign = [c for c in word if not c.isascii()]
    if len(foreign) != 1 or len(word) - 1 < 2:
        return False
    return unicodedata.name(foreign[0], "").startswith("LATIN")


def garbled(text: str) -> tuple[int, str] | None:
    """The garbled-text probe: (how many words look garbled, the commonest one), or None when
    there are too few to matter (GARBLED_MIN_WORDS, GARBLED_MIN_SHARE). Run on what extraction
    produced, after any repair — it reports what is still broken, whatever the cause."""
    from collections import Counter  # noqa: PLC0415

    total = 0
    bad: Counter = Counter()
    for match in _WORD.finditer(text):
        total += 1
        if _looks_garbled(match.group(0)):
            bad[match.group(0)] += 1
    count = sum(bad.values())
    if count < GARBLED_MIN_WORDS or count < GARBLED_MIN_SHARE * total:
        return None
    return count, bad.most_common(1)[0][0]


def probe(path: Path) -> Probe:
    """Cheap facts for the pre-flight report: slide and page counts, expected status."""
    extension = path.suffix.lower()
    if extension in MEDIA_EXTENSIONS:
        return Probe(status=MEDIA, reason="audio/video: recorded, content not extracted")
    extractor = EXTRACTORS.get(extension)
    if extractor is None:
        return Probe(status=UNSUPPORTED, reason=_unsupported_hint(extension))
    captured: list[str] = []
    try:
        with quiet(captured):
            found = extractor.probe(path)
        found.warnings = found.warnings + captured
        return found
    except Exception as exc:
        return Probe(status=UNSUPPORTED, reason=f"could not be read ({type(exc).__name__})")


def seconds_for(path: Path) -> float:
    extractor = EXTRACTORS.get(path.suffix.lower())
    return extractor.seconds if extractor else 0.0


# -- anchors -----------------------------------------------------------------

_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")
_FENCE = re.compile(r"^[ \t]*(```|~~~)")


def slug(text: str) -> str:
    """A heading's anchor: lowercased, every run of non-word characters → '-'.

    `Slide 18` → `slide-18`, `Question 3: heaps` → `question-3-heaps`. Unicode letters are kept,
    so a heading in another script still gets an anchor.
    """
    return re.sub(r"[\W_]+", "-", text.lower()).strip("-")


def headings(markdown: str) -> list[tuple[int, str]]:
    """Every ATX heading outside fenced code blocks: (level, text)."""
    found: list[tuple[int, str]] = []
    fenced = False
    for line in markdown.splitlines():
        if _FENCE.match(line):
            fenced = not fenced
            continue
        if fenced:
            continue
        match = _HEADING.match(line)
        if match:
            found.append((len(match.group(1)), match.group(2)))
    return found


def anchors(markdown: str) -> list[str]:
    """The anchors a locator can name in this document, in order.

    A repeated heading gets `-1`, `-2`, … appended, as GitHub's Markdown renderer does, so the
    anchor a teacher sees in a preview is the one a locator uses.
    """
    seen: dict[str, int] = {}
    result: list[str] = []
    for _level, text in headings(markdown):
        base = slug(text)
        if not base:
            continue
        if base in seen:
            seen[base] += 1
            result.append(f"{base}-{seen[base]}")
        else:
            seen[base] = 0
            result.append(base)
    return result


#: A label in an index is one line, and short: it names a place, it does not quote it.
LABEL_CHARS = 120


def label(text: str) -> str:
    """One index label: a single line, at most LABEL_CHARS, never a heading."""
    text = _clean(text)
    if len(text) > LABEL_CHARS:
        text = text[:LABEL_CHARS - 1].rstrip() + "…"
    return _escape_heading(text)


def index_body(markdown: str, labels: dict[str, list[str]]) -> str:
    """The index of a private material (D-040): every anchor heading the full text has, in the
    same order and so with the same anchors, each followed only by its one-line labels. No body
    text — this is what is committed for a file whose text may not be."""
    lines: list[str] = []
    seen: dict[str, int] = {}
    for level, text in headings(markdown):
        lines += [f"{'#' * level} {text}", ""]
        base = slug(text)
        if not base:
            continue
        if base in seen:
            seen[base] += 1
            anchor = f"{base}-{seen[base]}"
        else:
            seen[base] = 0
            anchor = base
        for line in labels.get(anchor, []):
            lines += [line, ""]
    return "\n".join(lines).rstrip() + "\n" if lines else ""


def _clean(text: str | None) -> str:
    return " ".join(str(text or "").split())


# Metadata titles that are really the authoring tool talking, not the author.
_JUNK_TITLE = re.compile(r"^(powerpoint presentation|presentation\d*|untitled.*|slide \d+|"
                         r"microsoft (word|powerpoint) - .*|document\d*)$", re.I)
# …or a file name the tool recorded (`manual.dvi`, `chapter3.tex`, `C:\\notes\\x.doc`) — F-12.
_FILE_NAME_TITLE = re.compile(r"^\S+\.(dvi|tex|ps|eps|pdf|docx?|pptx?|odt|odp|rtf|txt|md|html?|"
                              r"indd|qxd|key|pages)$", re.I)


def _title_from_metadata(value: str | None) -> str:
    value = _clean(value)
    if not value or _JUNK_TITLE.match(value) or _FILE_NAME_TITLE.match(value):
        return ""
    if "\\" in value or value.startswith("/"):  # a path, not a title
        return ""
    return value


# -- built in: Markdown and plain text ----------------------------------------

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


@register(".md", ".markdown", seconds=0.01)
def extract_markdown(path: Path) -> Extraction:
    text = _read_text(path).replace("\r\n", "\n")
    found = headings(text)
    title = found[0][1] if found else ""
    return Extraction(INGESTED, body=text.rstrip() + "\n", title=title)


@register(".txt", ".text", seconds=0.01)
def extract_text(path: Path) -> Extraction:
    text = _read_text(path).replace("\r\n", "\n")
    # A plain-text file has no headings, so it has no anchors: it is cited by its id alone.
    # A line that happens to start with '#' must not become one, so it is escaped.
    body = "\n".join(("\\" + line) if line.lstrip().startswith("#") else line
                     for line in text.splitlines())
    first = next((line.strip() for line in text.splitlines() if line.strip()), "")
    return Extraction(INGESTED, body=body.rstrip() + "\n", title=first[:80], title_from_body=True)


# -- built in: PowerPoint -----------------------------------------------------

def _pptx_probe(path: Path) -> Probe:
    import zipfile  # noqa: PLC0415

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        slides = sum(1 for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    return Probe(slides=slides)


@register(".pptx", probe=_pptx_probe, seconds=0.05)
def extract_pptx(path: Path) -> Extraction:
    from pptx import Presentation  # noqa: PLC0415

    deck = Presentation(str(path))
    lines: list[str] = []
    labels: dict[str, list[str]] = {}
    first_title = ""
    empty = 0

    for number, slide in enumerate(deck.slides, start=1):
        anchor = f"slide-{number}"
        lines += [f"## Slide {number}", ""]
        start = len(lines)
        if slide._element.get("show") == "0":
            lines += ["*(hidden slide)*", ""]
            labels.setdefault(anchor, []).append("*(hidden slide)*")

        title_shape = slide.shapes.title
        title = (_clean(" ".join(_pptx_paragraph_text(p) for p in title_shape.text_frame.paragraphs))
                 if title_shape is not None and title_shape.has_text_frame else "")
        if title:
            if number == 1:
                first_title = _title_from_metadata(title)  # the same boilerplate rejected
            lines += [f"**{title}**", ""]
            labels.setdefault(anchor, []).append(f"**{label(title)}**")

        for shape in _shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = _clean(_pptx_paragraph_text(paragraph))
                    if text:
                        lines.append(("  " * paragraph.level) + f"- {_escape_heading(text)}")
                lines.append("")
            if getattr(shape, "has_table", False) and shape.has_table:
                lines += _table_rows([[cell.text for cell in row.cells] for row in shape.table.rows])
                lines.append("")
            click = getattr(shape, "click_action", None)
            address = getattr(getattr(click, "hyperlink", None), "address", None) if click else None
            if address:
                lines += [f"*(link: {address})*", ""]

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append("> **Notes:** " + " ".join(notes.split()))
                lines.append("")

        if not any(line.strip() and line != "*(hidden slide)*" for line in lines[start:]):
            empty += 1

    # A deck's title is slide 1's title, else its file name (F-12, teacher test): never a later
    # slide's — a section, a digression or a tribute slide — and not the metadata title, which
    # is the authoring tool's or the template's as often as the author's.
    body = "\n".join(lines).rstrip() + "\n"
    return Extraction(INGESTED, body=body, title=first_title, labels=labels,
                      empty=(empty, len(deck.slides), "slides"),
                      low_yield=_low_yield(path, body, r"ppt/slides/slide\d+\.xml", "deck"))


def _pptx_paragraph_text(paragraph) -> str:
    """A slide paragraph's text, in order: runs (a hyperlink kept as a Markdown link, so its URL is
    readable in the ingested text — links are not harvested as materials, D-040), fields, line
    breaks, and Office Math equations (F-07), which python-pptx's `runs` leaves out."""
    from pptx.text.text import _Run  # noqa: PLC0415

    parts: list[str] = []
    for child in paragraph._p.iterchildren():
        tag = child.tag
        if tag == f"{_A}r":
            run = _Run(child, paragraph)
            parts.append(_linked(run.text, run.hyperlink.address if run.hyperlink is not None else None))
        elif tag == f"{_A}fld":
            parts.append("".join(t.text or "" for t in child.iter(f"{_A}t")))
        elif tag == f"{_A}br":
            parts.append(" ")
        elif tag == f"{_MC}AlternateContent":
            parts.append(_alternate_text(child))
    return "".join(parts)


def _alternate_text(alternate) -> str:
    """`mc:AlternateContent` in a slide: an equation in its `Choice` (PowerPoint's Office Math),
    else whatever text its `Fallback` has — never both, they are the same content twice."""
    choice = alternate.find(f"{_MC}Choice")
    if choice is not None:
        math = next((e for e in choice.iter() if e.tag in (f"{_M}oMathPara", f"{_M}oMath")), None)
        if math is not None:
            return f" {equation(math)} "
        text = "".join(t.text or "" for t in choice.iter(f"{_A}t"))
        if text:
            return text
    fallback = alternate.find(f"{_MC}Fallback")
    return "".join(t.text or "" for t in fallback.iter(f"{_A}t")) if fallback is not None else ""


def _low_yield(path: Path, body: str, parts: str, what: str) -> str:
    """The low-yield note (F-06), or "": text far smaller than the document XML it came from."""
    import zipfile  # noqa: PLC0415

    try:
        with zipfile.ZipFile(path) as archive:
            xml = sum(info.file_size for info in archive.infolist() if re.fullmatch(parts, info.filename))
    except (OSError, zipfile.BadZipFile):
        return ""
    text = "".join(line for line in body.splitlines()
                   if not line.startswith("## ") and line.strip() != "*(hidden slide)*")
    characters = len("".join(text.split()))
    if xml < LOW_YIELD_MIN_XML or characters >= LOW_YIELD_RATIO * xml:
        return ""
    return (f"only {characters} characters of text from a {_kb(path.stat().st_size)} {what} — text "
            "boxes, pictures or embedded objects may not have been read; check the ingested file")


def _kb(size: int) -> str:
    return f"{size / 1024:.0f} KB" if size < 1024 * 1024 else f"{size / 1024 / 1024:.1f} MB"


# -- Office Math (OMML) as linear text — F-07 ----------------------------------

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def equation(element) -> str:
    """An Office Math equation (`m:oMath`, `m:oMathPara`) as linear text in a code span —
    `(a)/(b)`, `x^(2)`, `∑_(i=1)^(n) i` — or `[equation]` when nothing can be read from it, so a
    reader knows something is there. For an algorithms course the equations are the content."""
    text = _clean(_omml(element))
    return f"`{text.replace('`', chr(39))}`" if text else "[equation]"


def _omml(element) -> str:
    tag = element.tag
    if not isinstance(tag, str) or not tag.startswith(_M):
        # Office Math may hold a word-processing run (w:r → w:t) or a drawing run (a:t)
        if tag in (f"{_W}t", f"{_A}t"):
            return element.text or ""
        return "".join(_omml(child) for child in element)
    name = tag[len(_M):]

    def part(child: str) -> str:
        found = element.find(f"{_M}{child}")
        return _clean(_omml(found)) if found is not None else ""

    def prop(path: str, default: str) -> str:
        found = element.find(path)
        return found.get(f"{_M}val", default) if found is not None else default

    if name == "t":
        return element.text or ""
    if name.endswith("Pr") or name == "ctrlPr":
        return ""
    if name == "f":
        return f"({part('num')})/({part('den')})"
    if name == "sSup":
        return f"{part('e')}^({part('sup')})"
    if name == "sSub":
        return f"{part('e')}_({part('sub')})"
    if name == "sSubSup":
        return f"{part('e')}_({part('sub')})^({part('sup')})"
    if name == "sPre":
        return f"_({part('sub')})^({part('sup')}){part('e')}"
    if name == "rad":
        degree = part("deg")
        return f"root({degree})({part('e')})" if degree else f"√({part('e')})"
    if name == "nary":
        symbol = prop(f"{_M}naryPr/{_M}chr", "∫")
        low, high = part("sub"), part("sup")
        return symbol + (f"_({low})" if low else "") + (f"^({high})" if high else "") + " " + part("e")
    if name == "d":
        begin = prop(f"{_M}dPr/{_M}begChr", "(")
        end = prop(f"{_M}dPr/{_M}endChr", ")")
        separator = prop(f"{_M}dPr/{_M}sepChr", "|")
        return begin + f" {separator} ".join(_clean(_omml(e)) for e in element.findall(f"{_M}e")) + end
    if name == "func":
        return f"{part('fName')} {part('e')}"
    if name in ("limLow", "limUpp"):
        return f"{part('e')}{'_' if name == 'limLow' else '^'}({part('lim')})"
    if name in ("acc", "bar"):
        mark = prop(f"{_M}accPr/{_M}chr", "̂") if name == "acc" else "‾"
        return f"{part('e')}{mark}"
    if name == "m":
        rows = ["; ".join(_clean(_omml(e)) for e in row.findall(f"{_M}e")) for row in element.findall(f"{_M}mr")]
        return "[" + " | ".join(rows) + "]"
    if name in ("eqArr", "oMathPara"):
        return "; ".join(_clean(_omml(child)) for child in element if not child.tag.endswith("Pr"))
    return "".join(_omml(child) for child in element)


def _linked(text: str, address: str | None) -> str:
    """A run of text, as a Markdown link when it carries a hyperlink."""
    if not address or not text.strip():
        return text
    return f"[{text}]({address})"


def _shapes(shapes):
    """Every shape, descending into groups, in slide order."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _shapes(shape.shapes)
        else:
            yield shape


# -- built in: PDF (text layer) -----------------------------------------------

def _pdf_probe(path: Path) -> Probe:
    from pypdf import PdfReader  # noqa: PLC0415

    reader = PdfReader(str(path))
    return Probe(pages=len(reader.pages))


@register(".pdf", probe=_pdf_probe, seconds=0.08)
def extract_pdf(path: Path) -> Extraction:
    from pypdf import PdfReader  # noqa: PLC0415

    reader = PdfReader(str(path))
    lines: list[str] = []
    characters = 0
    printed = list(getattr(reader, "page_labels", []) or [])
    sections = _outline_by_page(reader)
    labels: dict[str, list[str]] = {}

    empty = 0
    for index, page in enumerate(reader.pages):
        number = index + 1
        anchor = f"page-{number}"
        lines += [f"## Page {number}", ""]
        # Locators use the physical page, which always exists and never repeats. A printed
        # page number that differs (front matter, a textbook's own numbering) is noted, so a
        # teacher citing "p. 45" can find page-63.
        page_label = printed[index] if index < len(printed) else str(number)
        if page_label and page_label != str(number):
            lines += [f"*(printed page {page_label})*", ""]
            labels.setdefault(anchor, []).append(f"*(printed page {label(page_label)})*")
        for title in sections.get(index, []):
            labels.setdefault(anchor, []).append(f"*(section: {label(title.translate(LIGATURES))})*")
        try:
            text = repair_ligatures((page.extract_text() or "").translate(LIGATURES))
        except Exception:  # one bad page must not lose the rest of the document
            text = ""
        page_characters = len("".join(text.split()))
        characters += page_characters
        empty += page_characters == 0
        for line in text.splitlines():
            line = line.rstrip()
            lines.append(_escape_heading(line) if line else "")
        lines.append("")
        # A link annotation's target is not in the text layer; it is noted, so the URL stays
        # readable in the ingested text (links are not harvested as materials, D-040).
        for uri in _link_annotations(page):
            if uri not in text:
                lines += [f"*(link: {uri})*", ""]

    pages = len(reader.pages)
    # A PDF's title is its metadata title, else the file name — never its first line of text,
    # which is a running head, a copyright line or a page number as often as a title (F-12).
    metadata_title = _title_from_metadata((reader.metadata or {}).get("/Title"))
    body = "\n".join(lines).rstrip() + "\n"

    if pages == 0 or characters < NO_TEXT_CHARS_PER_PAGE * pages:
        # Scanned: no text to extract. The page headings are still written, so a locator to a
        # page that exists resolves, and the teacher may type in the text by hand.
        return Extraction(NO_TEXT, body=body, title=metadata_title,
                          reason="no text layer (a scan?) — OCR is not done in Core; "
                                 "pages are anchored, the text is empty", labels=labels)

    return Extraction(INGESTED, body=body, title=metadata_title, labels=labels,
                      empty=(empty, pages, "pages"))


def _link_annotations(page) -> list[str]:
    """The URLs of a page's link annotations, in order, each once. A broken annotation is skipped."""
    found: list[str] = []
    try:
        for annotation in page.get("/Annots") or []:
            obj = annotation.get_object()
            action = obj.get("/A") or {}
            if obj.get("/Subtype") == "/Link" and "/URI" in action:
                uri = str(action["/URI"]).strip()
                if uri and uri not in found:
                    found.append(uri)
    except Exception:  # a malformed annotation must not lose the page
        pass
    return found


def _outline_by_page(reader) -> dict[int, list[str]]:
    """The PDF's outline (bookmarks) as page index → titles of the sections starting there, in
    outline order. Empty for a PDF without one — a scan, some exports — or a broken one: the
    index then has pages and printed labels only."""
    found: dict[int, list[str]] = {}

    def walk(items) -> None:
        for item in items:
            if isinstance(item, list):
                walk(item)
                continue
            try:
                page = reader.get_destination_page_number(item)
            except Exception:  # a dangling or malformed bookmark names no page
                continue
            title = _clean(getattr(item, "title", None) or "")
            if title and page is not None and page >= 0:
                found.setdefault(page, []).append(title)

    try:
        walk(reader.outline or [])
    except Exception:
        return {}
    return found


# -- built in: Word -----------------------------------------------------------

@register(".docx", seconds=0.2)
def extract_docx(path: Path) -> Extraction:
    import docx  # noqa: PLC0415
    from docx.oxml.ns import qn  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    document = docx.Document(str(path))
    lines: list[str] = []
    headings_found: list[str] = []

    def blocks(container) -> None:
        """Paragraphs and tables in order — and after each paragraph, the text boxes anchored in
        it (F-06: an official syllabus made of text boxes came out as 13 characters)."""
        for child in container.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, document)
                text = _clean(_docx_paragraph_text(child, paragraph))
                if text:
                    level = _heading_level(paragraph.style.name if paragraph.style is not None else "")
                    if level:
                        headings_found.append(text)
                        lines.extend(["", f"{'#' * level} {text}", ""])
                    elif _bold_label(child, paragraph, text):
                        # A form with no heading styles (BGU's syllabus: "Assessment:") still gets
                        # anchors. Not a title candidate: a label names a section, not the document.
                        lines.extend(["", f"## {text.rstrip(':').rstrip()}", ""])
                    elif paragraph.style is not None and "List" in paragraph.style.name:
                        lines.append(f"- {_escape_heading(text)}")
                    else:
                        lines.extend([_escape_heading(text), ""])
                for box in _text_boxes(child):
                    blocks(box)
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                lines.extend([""] + _table_rows([[cell.text for cell in row.cells] for row in table.rows]) + [""])

    blocks(document.element.body)
    markdown = "\n".join(lines).strip() + "\n"
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)

    title = (headings_found[0] if headings_found else "") or _title_from_metadata(document.core_properties.title)
    return Extraction(INGESTED, body=markdown, title=title,
                      low_yield=_low_yield(path, markdown, r"word/document\.xml", "document"))


def _docx_paragraph_text(element, paragraph) -> str:
    """A paragraph's text, in order: runs; hyperlinks kept as Markdown links (D-040: a link inside
    a document is readable in its ingested text, not harvested as a material); Office Math as
    linear text (F-07); and the runs inside tracked insertions, smart tags and content controls.
    Text boxes are not here — `_text_boxes` reads them, once."""
    from docx.text.hyperlink import Hyperlink  # noqa: PLC0415
    from docx.text.run import Run  # noqa: PLC0415

    parts: list[str] = []
    for child in element.iterchildren():
        tag = child.tag
        if tag == f"{_W}r":
            parts.append(Run(child, paragraph).text)
        elif tag == f"{_W}hyperlink":
            link = Hyperlink(child, paragraph)
            parts.append(_linked(link.text, link.address))
        elif tag in (f"{_M}oMath", f"{_M}oMathPara"):
            parts.append(f" {equation(child)} ")
        elif tag in (f"{_W}ins", f"{_W}smartTag", f"{_W}sdt", f"{_W}sdtContent", f"{_W}fldSimple",
                     f"{_W}customXml"):
            parts.append(_docx_paragraph_text(child, paragraph))
    return "".join(parts)


def _text_boxes(element) -> list:
    """The text boxes (`w:txbxContent`) anchored in a paragraph — the outermost ones only (a box
    in a box is read with its parent), and from `mc:AlternateContent` the `Choice` or the
    `Fallback`, never both: Word writes the same box twice, as DrawingML and as VML."""
    boxes = []
    for box in element.iter(f"{_W}txbxContent"):
        ancestor, keep = box.getparent(), True
        while ancestor is not None and ancestor is not element:
            if ancestor.tag == f"{_W}txbxContent":
                keep = False
                break
            if ancestor.tag == f"{_MC}Fallback":
                choice = ancestor.getparent().find(f"{_MC}Choice")
                if choice is not None and next(choice.iter(f"{_W}txbxContent"), None) is not None:
                    keep = False
                    break
            ancestor = ancestor.getparent()
        if keep:
            boxes.append(box)
    return boxes


#: A bold label read as a heading is at most this long (teacher test: "Learning outcomes of the module:").
BOLD_LABEL_CHARS = 80


def _bold_label(element, paragraph, text: str) -> bool:
    """A short paragraph, entirely bold, ending in ':' — a section label in a document without
    heading styles, read as a `##` heading so it can be cited (`M0017#assessment`)."""
    from docx.text.run import Run  # noqa: PLC0415

    if len(text) > BOLD_LABEL_CHARS or not text.endswith(":") or not text.rstrip(":").strip():
        return False
    style_bold = bool(paragraph.style is not None and paragraph.style.font.bold)
    runs = [Run(r, paragraph) for r in element.iter(f"{_W}r")]
    runs = [run for run in runs if run.text.strip()]
    return bool(runs) and all(run.bold or (run.bold is None and style_bold) for run in runs)


def _heading_level(style: str) -> int:
    if style == "Title":
        return 1
    match = re.fullmatch(r"Heading (\d)", style)
    return min(int(match.group(1)), 6) if match else 0


# -- optional: pandoc and LibreOffice, used only when installed ---------------

PANDOC_EXTENSIONS = {".odt": "odt", ".rtf": "rtf", ".html": "html", ".htm": "html", ".epub": "epub"}
LIBREOFFICE_EXTENSIONS = {".ppt": "pptx", ".odp": "pptx", ".doc": "docx"}


def pandoc() -> str | None:
    return shutil.which("pandoc")


def libreoffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def _unsupported_hint(extension: str) -> str:
    if extension in PANDOC_EXTENSIONS:
        return "install pandoc to read this format, or export it to PDF"
    if extension in LIBREOFFICE_EXTENSIONS:
        return "install LibreOffice to read this format, or export it to PDF"
    if extension in {".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".tif", ".tiff", ".bmp", ".heic"}:
        return "an image: text in images is not extracted in Core (no OCR)"
    if extension in {".zip", ".tar", ".gz", ".7z", ".rar"}:
        return "an archive: unpack it into materials/source/ and re-run"
    if extension in {".xlsx", ".xls", ".csv", ".ods"}:
        return "a spreadsheet: no extractor in Core — export the relevant part to PDF"
    return "no extractor for this format — export it to PDF and re-run"


def _optional_probe(path: Path) -> Probe:
    extension = path.suffix.lower()
    if extension in PANDOC_EXTENSIONS and pandoc() is None:
        return Probe(status=UNSUPPORTED, reason=_unsupported_hint(extension))
    if extension in LIBREOFFICE_EXTENSIONS and libreoffice() is None:
        return Probe(status=UNSUPPORTED, reason=_unsupported_hint(extension))
    return Probe()


@register(*PANDOC_EXTENSIONS, probe=_optional_probe, seconds=1.0)
def extract_with_pandoc(path: Path) -> Extraction:
    tool = pandoc()
    if tool is None:
        return Extraction(UNSUPPORTED, reason=_unsupported_hint(path.suffix.lower()))
    source_format = PANDOC_EXTENSIONS[path.suffix.lower()]
    completed = subprocess.run(
        [tool, "--from", source_format, "--to", "gfm", "--wrap=none", str(path)],
        capture_output=True, text=True, timeout=300, check=True,
    )
    markdown = completed.stdout
    found = headings(markdown)
    return Extraction(INGESTED, body=markdown.rstrip() + "\n", title=found[0][1] if found else "")


@register(*LIBREOFFICE_EXTENSIONS, probe=_optional_probe, seconds=4.0)
def extract_with_libreoffice(path: Path) -> Extraction:
    tool = libreoffice()
    if tool is None:
        return Extraction(UNSUPPORTED, reason=_unsupported_hint(path.suffix.lower()))
    target = LIBREOFFICE_EXTENSIONS[path.suffix.lower()]
    # Converted into a scratch directory — never next to the teacher's file in source/.
    with tempfile.TemporaryDirectory(prefix="classkit-") as scratch:
        subprocess.run(
            [tool, "--headless", "--convert-to", target, "--outdir", scratch, str(path)],
            capture_output=True, timeout=300, check=True,
        )
        converted = Path(scratch) / f"{path.stem}.{target}"
        return EXTRACTORS[f".{target}"].extract(converted)


# -- helpers ------------------------------------------------------------------

def _escape_heading(text: str) -> str:
    """Body text that starts with '#' must not become a heading — it would invent an anchor."""
    return ("\\" + text) if text.lstrip().startswith("#") else text


def _table_rows(rows: list[list[str]]) -> list[str]:
    if not rows:
        return []
    cleaned = [[_clean(cell).replace("|", "\\|") for cell in row] for row in rows]
    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    out = ["| " + " | ".join(cleaned[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(row) + " |" for row in cleaned[1:]]
    return out

